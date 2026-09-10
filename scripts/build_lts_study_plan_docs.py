#!/usr/bin/env python3
"""Build branded Word (.docx) study plan for London Tamil Sangam Class 3."""
from __future__ import annotations

from pathlib import Path

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import nsmap, qn
from docx.shared import Cm, Emu, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path("/workspace")
HTML = ROOT / "mundram-vaguppu-paadathittam.html"
OUT_DOCX = ROOT / "LTS_Class3_Tamil_Study_Plan_2025-26.docx"
CREST = ROOT / "scripts" / "lts_crest.png"

SANGAM = RGBColor(0x7A, 0x18, 0x36)
NAVY = RGBColor(0x1B, 0x2A, 0x4A)
GOLD = RGBColor(0xC4, 0xA3, 0x5A)
INK = RGBColor(0x1D, 0x20, 0x33)
SOFT = RGBColor(0x5A, 0x61, 0x78)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CREAM = "F8F3EA"
TYPE_COLOURS = {
    "song": "7B2D8E",
    "story": "0E7C66",
    "talk": "C25E00",
    "word": "0B6FA4",
    "gram": "4A3AA8",
    "value": "B3123F",
}
FONT_SANS = "Noto Sans Tamil"
FONT_SERIF = "Noto Serif Tamil"
TAMIL_BOLD = "/usr/share/fonts/truetype/noto/NotoSerifTamil-Bold.ttf"
LATIN_BOLD = "/usr/share/fonts/truetype/noto/NotoSansTamil-Bold.ttf"


def hex_rgb(h: str) -> RGBColor:
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def shade_cell(cell, fill: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_borders(cell, color="C4A35A", sz="8") -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def set_table_borders(table, color="E4D9C8", sz="4") -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)


def no_table_borders(table) -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def prevent_row_split(row) -> None:
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    trPr.append(cant)


def set_run_font(run, name=FONT_SANS, size=11, bold=False, italic=False, color=INK) -> None:
    run.font.name = name
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:cs"), name)
    rFonts.set(qn("w:eastAsia"), name)
    sz = str(int(size * 2))
    for tag in ("w:sz", "w:szCs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            rPr.append(el)
        el.set(qn("w:val"), sz)
    if bold:
        bCs = rPr.find(qn("w:bCs"))
        if bCs is None:
            bCs = OxmlElement("w:bCs")
            rPr.append(bCs)
    lang = rPr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rPr.append(lang)
    lang.set(qn("w:val"), "en-GB")
    lang.set(qn("w:bidi"), "ta-IN")


def add_run(p, text, **kwargs):
    if text is None:
        return None
    run = p.add_run(str(text))
    set_run_font(run, **kwargs)
    return run


def para(doc, text="", *, style=None, space_after=8, space_before=0, align=None, keep=False):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    if align:
        p.alignment = align
    if keep:
        p.paragraph_format.keep_with_next = True
    if text:
        add_run(p, text)
    return p


def heading(doc, text, level=1, english=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    colour = SANGAM if level == 1 else NAVY
    size = 18 if level == 1 else 14 if level == 2 else 12
    add_run(p, text, name=FONT_SERIF, size=size, bold=True, color=colour)
    if english:
        add_run(p, "  ", name=FONT_SANS, size=10, color=SOFT)
        add_run(p, english, name=FONT_SANS, size=10, italic=True, color=SOFT)
    if level == 1:
        # gold rule under H1
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "12")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), "C4A35A")
        pBdr.append(bottom)
        pPr.append(pBdr)
    return p


def add_fld(paragraph, instr: str) -> None:
    run = paragraph.add_run()
    r = run._r
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instr
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r.append(begin)
    r.append(text)
    r.append(sep)
    r.append(end)
    set_run_font(run, name=FONT_SANS, size=9, color=SOFT)


def add_html_runs(p, node, *, size=11, color=INK):
    if node is None:
        return
    for child in node.children:
        if isinstance(child, NavigableString):
            t = str(child)
            if t:
                add_run(p, t, size=size, color=color)
        elif isinstance(child, Tag):
            if child.name in ("b", "strong"):
                add_run(p, child.get_text(), size=size, bold=True, color=color)
            elif child.name in ("em", "i"):
                add_run(p, child.get_text(), size=size, italic=True, color=SOFT)
            elif child.name == "br":
                add_run(p, "\n", size=size, color=color)
            else:
                add_html_runs(p, child, size=size, color=color)


def make_crest(path: Path) -> None:
    img = Image.new("RGBA", (360, 360), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    d.ellipse((8, 8, 352, 352), fill=(122, 24, 54, 255))
    d.ellipse((28, 28, 332, 332), outline=(196, 163, 90, 255), width=8)
    d.ellipse((48, 48, 312, 312), outline=(232, 213, 163, 180), width=2)
    tamil = ImageFont.truetype(TAMIL_BOLD, 88)
    latin = ImageFont.truetype(LATIN_BOLD, 36)
    yearf = ImageFont.truetype(LATIN_BOLD, 28)
    zha = "ழ"
    bbox = d.textbbox((0, 0), zha, font=tamil)
    d.text(((360 - (bbox[2] - bbox[0])) / 2, 86), zha, font=tamil, fill=(232, 213, 163, 255))
    for label, y, font in (("LTS", 210, latin), ("1936", 262, yearf)):
        bbox = d.textbbox((0, 0), label, font=font)
        d.text(((360 - (bbox[2] - bbox[0])) / 2, y), label, font=font, fill=(255, 255, 255, 255) if label == "LTS" else (232, 213, 163, 255))
    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path)


def colour_from_style(style: str) -> str:
    if not style:
        return TYPE_COLOURS["gram"]
    for key, val in TYPE_COLOURS.items():
        if f"var(--{key})" in style or f"--{key}" in style:
            return val
    return TYPE_COLOURS["gram"]


def set_narrow_margins(section) -> None:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(1.8)
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.7)


def add_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    section.different_first_page_header_footer = True

    # First page: empty header (cover has its own letterhead)
    section.first_page_header.paragraphs[0].text = ""
    fp_f = section.first_page_footer.paragraphs[0]
    fp_f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(fp_f, "The London Tamil Sangam · Charity no. 1097724 · www.ltsuk.org", size=8, color=SOFT)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_run(header, "இலண்டன் தமிழ் சங்கம்  ·  The London Tamil Sangam", name=FONT_SERIF, size=9, bold=True, color=SANGAM)
    add_run(header, "    Class 3 Tamil Study Plan  2025–2026", name=FONT_SANS, size=9, color=NAVY)
    pPr = header._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "18")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "7A1836")
    pBdr.append(bottom)
    pPr.append(pBdr)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_run(footer, "Charity 1097724  ·  369 High Street North, Manor Park, London E12 6PG  ·  ", size=8, color=SOFT)
    add_fld(footer, " PAGE ")
    add_run(footer, "  |  Tamil Language School", size=8, color=SOFT)


def banner_table(doc, left_title, left_sub, fill="7A1836") -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_borders(cell, color=fill, sz="0")
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    add_run(p, left_title, name=FONT_SERIF, size=16, bold=True, color=WHITE)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(8)
    add_run(p2, left_sub, name=FONT_SANS, size=10, color=GOLD)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)


def add_card_row(doc, cards: list[tuple[str, str, str]]) -> None:
    table = doc.add_table(rows=1, cols=len(cards))
    table.autofit = True
    set_table_borders(table, color="E4D9C8", sz="6")
    for i, (title, english, body) in enumerate(cards):
        cell = table.cell(0, i)
        shade_cell(cell, "FFFCF6")
        p = cell.paragraphs[0]
        add_run(p, title, name=FONT_SERIF, size=12, bold=True, color=SANGAM)
        p2 = cell.add_paragraph()
        add_run(p2, english, name=FONT_SANS, size=9, italic=True, color=SOFT)
        p3 = cell.add_paragraph()
        p3.paragraph_format.space_after = Pt(4)
        add_run(p3, body, size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def html_table_to_docx(doc, table_el: Tag, header_fill="1B2A4A") -> None:
    rows = table_el.find_all("tr")
    if not rows:
        return
    cols = max(len(r.find_all(["th", "td"])) for r in rows)
    wt = doc.add_table(rows=len(rows), cols=cols)
    wt.autofit = True
    set_table_borders(table=wt, color="D7CBB8", sz="4")
    for r_i, tr in enumerate(rows):
        cells = tr.find_all(["th", "td"])
        prevent_row_split(wt.rows[r_i])
        for c_i in range(cols):
            cell = wt.cell(r_i, c_i)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            if c_i < len(cells):
                src = cells[c_i]
                is_h = src.name == "th" or r_i == 0
                if is_h:
                    shade_cell(cell, header_fill)
                    add_run(p, src.get_text(" ", strip=True), size=9, bold=True, color=WHITE)
                else:
                    shade_cell(cell, "FFFCF6")
                    add_html_runs(p, src, size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def add_list(doc, el: Tag, numbered=False) -> None:
    items = el.find_all("li", recursive=False)
    for i, li in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_after = Pt(3)
        prefix = f"{i}.  " if numbered else "•  "
        add_run(p, prefix, size=11, bold=numbered, color=SANGAM if numbered else NAVY)
        add_html_runs(p, li, size=11)


def add_block(doc, block: Tag, accent="7A1836") -> None:
    label = block.select_one(".blabel")
    if label:
        p = para(doc, space_after=3, space_before=8, keep=True)
        add_run(p, label.get_text(" ", strip=True), name=FONT_SERIF, size=11, bold=True, color=hex_rgb(accent))
    for child in block.children:
        if not isinstance(child, Tag):
            continue
        if "blabel" in child.get("class", []):
            continue
        if child.name == "ul":
            add_list(doc, child, numbered=False)
        elif child.name == "ol":
            add_list(doc, child, numbered=True)
        elif child.name == "table":
            html_table_to_docx(doc, child)
        elif child.name == "div" and "chips" in child.get("class", []):
            p = para(doc, space_after=6)
            chips = [c.get_text(strip=True) for c in child.select(".chip")]
            add_run(p, "  ·  ".join(chips), size=11, color=NAVY)
        elif child.name == "div" and any(c in child.get("class", []) for c in ("hw", "note", "game")):
            box = doc.add_table(rows=1, cols=1)
            set_table_borders(box, color=accent, sz="8")
            cell = box.cell(0, 0)
            shade_cell(cell, "FBF6EE")
            p = cell.paragraphs[0]
            add_html_runs(p, child, size=11)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
        elif child.name in ("p", "div"):
            p = para(doc, space_after=6)
            add_html_runs(p, child, size=11)
        elif child.name == "span" and "mini" in child.get("class", []):
            p = para(doc, space_after=4)
            add_run(p, child.get_text(" ", strip=True), size=10, italic=True, color=SOFT)


def add_lesson(doc, article: Tag) -> None:
    style = article.get("style", "")
    fill = colour_from_style(style)
    num = article.select_one(".lhead .num")
    title = article.select_one(".lhead h3")
    typ = article.select_one(".lhead .type")
    head = "  ·  ".join(
        t.get_text(" ", strip=True) for t in (num, title, typ) if t
    )
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_borders(cell, color=fill, sz="0")
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    add_run(p, head, name=FONT_SERIF, size=13, bold=True, color=WHITE)

    for block in article.select(".block, .game"):
        add_block(doc, block, accent=fill)
    # leftover tables inside lesson but outside .block
    for tbl in article.select(".lmain > table, .lside > table"):
        html_table_to_docx(doc, tbl)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(10)


def add_year_map(doc, soup: BeautifulSoup) -> None:
    heading(doc, "ஆண்டு நீண்டகாலத் திட்டம்", 1, "Long-term plan — 30 weekly sessions")
    p = para(doc, space_after=10)
    add_run(
        p,
        "Use this as the scheme of work. Colour shows the curriculum strand. Move weeks around holidays as needed.",
        size=11,
        color=SOFT,
    )
    tiles = []
    section = soup.select_one("#year-map")
    if not section:
        return
    current_term = ""
    for el in section.select(".map-term, a.tile"):
        if "map-term" in el.get("class", []):
            current_term = el.get_text(" ", strip=True)
            tiles.append(("term", current_term, "", ""))
        else:
            tiles.append(
                (
                    "tile",
                    el.select_one(".wk").get_text(strip=True) if el.select_one(".wk") else "",
                    el.select_one(".nm").get_text(strip=True) if el.select_one(".nm") else "",
                    el.select_one(".tg").get_text(strip=True) if el.select_one(".tg") else "",
                )
            )
    wt = doc.add_table(rows=1, cols=3)
    hdr = wt.rows[0].cells
    for i, label in enumerate(("Week", "Lesson", "Strand")):
        shade_cell(hdr[i], "1B2A4A")
        p = hdr[i].paragraphs[0]
        add_run(p, label, size=10, bold=True, color=WHITE)
    for kind, a, b, c in tiles:
        row = wt.add_row()
        prevent_row_split(row)
        if kind == "term":
            merged = row.cells[0].merge(row.cells[2])
            shade_cell(merged, "7A1836")
            add_run(merged.paragraphs[0], a, size=10, bold=True, color=WHITE)
        else:
            shade_cell(row.cells[0], "FFFCF6")
            shade_cell(row.cells[1], "FFFCF6")
            shade_cell(row.cells[2], "FFFCF6")
            add_run(row.cells[0].paragraphs[0], a, size=10, bold=True, color=NAVY)
            add_run(row.cells[1].paragraphs[0], b, name=FONT_SERIF, size=11, color=INK)
            add_run(row.cells[2].paragraphs[0], c, size=10, color=SANGAM)
    set_table_borders(wt, color="E4D9C8", sz="4")
    doc.add_paragraph().paragraph_format.space_after = Pt(8)


def cover_page(doc: Document) -> None:
    # Maroon identity bar
    bar = doc.add_table(rows=1, cols=1)
    cell = bar.cell(0, 0)
    shade_cell(cell, "7A1836")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    add_run(p, "THE LONDON TAMIL SANGAM  ·  இலண்டன் தமிழ் சங்கம்", name=FONT_SERIF, size=11, bold=True, color=GOLD)
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(10)
    add_run(p2, "Tamil Language School  ·  தமிழ்ப் பள்ளி  ·  Founded 1936", size=10, color=WHITE)

    # Crest + titles
    ident = doc.add_table(rows=1, cols=2)
    no_table_borders(ident)
    ident.columns[0].width = Cm(3.4)
    ident.columns[1].width = Cm(14.2)
    left, right = ident.cell(0, 0), ident.cell(0, 1)
    left.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if CREST.exists():
        run = left.paragraphs[0].add_run()
        run.add_picture(str(CREST), width=Cm(2.8))
    p = right.paragraphs[0]
    p.paragraph_format.space_before = Pt(8)
    add_run(p, "இலண்டன் தமிழ் சங்கம்", name=FONT_SERIF, size=22, bold=True, color=SANGAM)
    p = right.add_paragraph()
    add_run(p, "The London Tamil Sangam", name=FONT_SERIF, size=14, bold=True, color=NAVY)
    p = right.add_paragraph()
    add_run(p, "Tamil Language School  ·  Complementary education for ages 5–16", size=10, color=SOFT)

    gold = doc.add_paragraph()
    pPr = gold._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "24")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "C4A35A")
    pBdr.append(bottom)
    pPr.append(pBdr)

    t = para(doc, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=16, space_after=4)
    add_run(t, "EDUCATION STUDY PLAN  ·  கல்வித் திட்டம்", size=10, bold=True, color=GOLD)

    t = para(doc, space_after=4)
    add_run(t, "மூன்றாம் வகுப்பு", name=FONT_SERIF, size=26, bold=True, color=NAVY)
    t = para(doc, space_after=4)
    add_run(t, "பாடத்திட்டம் · கற்றல் திட்டம்", name=FONT_SERIF, size=20, bold=True, color=SANGAM)
    t = para(doc, space_after=12)
    add_run(t, "Class 3 Tamil Study Plan & Scheme of Work", name=FONT_SERIF, size=16, italic=True, color=NAVY)

    intro = para(doc, space_after=14)
    add_run(
        intro,
        "A 30-week complementary-school curriculum for London Tamil Sangam pupils: "
        "24 lessons, 8 grammar units and three assessment points. Each weekly session is planned "
        "for 90 minutes, with learning outcomes, vocabulary, classroom activities, homework and assessment.",
        size=11,
        color=INK,
    )

    meta = doc.add_table(rows=5, cols=2)
    set_table_borders(meta, color="C4A35A", sz="6")
    rows = [
        ("Academic year", "2025–2026"),
        ("Document", "Class 3 Scheme of Work  ·  Edition 1"),
        ("Audience", "Teachers, parents and pupils"),
        ("Centres", "East Ham (Plashet School)  ·  Redbridge (Avanti Court)  ·  Online"),
        ("Pathway", "Cambridge GCE O Level Tamil"),
    ]
    for i, (k, v) in enumerate(rows):
        shade_cell(meta.cell(i, 0), "7A1836")
        add_run(meta.cell(i, 0).paragraphs[0], k, size=10, bold=True, color=WHITE)
        shade_cell(meta.cell(i, 1), "FFFCF6")
        add_run(meta.cell(i, 1).paragraphs[0], v, size=11, color=NAVY)

    p = para(doc, space_before=16, space_after=2)
    add_run(p, "School office  ·  020 8471 7672  ·  info.school@ltsuk.org  ·  www.ltsuk.org", size=10, color=SOFT)
    p = para(doc, space_after=2)
    add_run(p, "369 High Street North, Manor Park, London E12 6PG  ·  Registered charity 1097724", size=10, color=SOFT)
    doc.add_page_break()


def configure_normal(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = FONT_SANS
    style.font.size = Pt(11)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), FONT_SANS)
    rFonts.set(qn("w:hAnsi"), FONT_SANS)
    rFonts.set(qn("w:cs"), FONT_SANS)


def add_generic_section(doc, section: Tag, h_level=1) -> None:
    head = section.select_one(".sec-head")
    if head:
        ta = head.find("h2")
        en = head.select_one(".en")
        heading(
            doc,
            ta.get_text(" ", strip=True) if ta else "",
            h_level,
            en.get_text(" ", strip=True) if en else "",
        )
    intro = section.select_one(".sec-intro")
    if intro:
        p = para(doc, space_after=10)
        add_html_runs(p, intro, size=11, color=SOFT)
    # cards
    cards = []
    for card in section.select(".card"):
        h3 = card.find("h3")
        title = ""
        english = ""
        if h3:
            en = h3.select_one(".en")
            if en:
                english = en.get_text(" ", strip=True)
                title = h3.get_text(" ", strip=True).replace(english, "").strip()
            else:
                title = h3.get_text(" ", strip=True)
        body = " ".join(p.get_text(" ", strip=True) for p in card.find_all(["p", "li"]))
        cards.append((title, english, body))
    if cards:
        # split into rows of 3
        for i in range(0, len(cards), 3):
            add_card_row(doc, cards[i : i + 3])
    for lg in section.select(".lg"):
        p = para(doc, space_after=4)
        b = lg.find("b")
        add_run(p, (b.get_text(strip=True) + "  ") if b else "", name=FONT_SERIF, size=12, bold=True, color=SANGAM)
        span = lg.find("span")
        if span:
            add_run(p, span.get_text(" ", strip=True), size=11, color=SOFT)
    for rule in section.select(".rule"):
        h = rule.find("h4")
        if h:
            heading(doc, h.get_text(" ", strip=True), 3)
        ul = rule.find("ul")
        if ul:
            add_list(doc, ul)
    for table in section.find_all("table"):
        # skip tables inside lessons
        if table.find_parent("article"):
            continue
        html_table_to_docx(doc, table)
    # pct bar caption
    pct = section.select_one(".pct")
    if pct:
        p = para(doc, space_after=8)
        add_run(p, pct.get_text(" ", strip=True), size=11, bold=True, color=NAVY)


def main() -> None:
    make_crest(CREST)
    soup = BeautifulSoup(HTML.read_text(encoding="utf-8"), "lxml")
    doc = Document()
    configure_normal(doc)
    set_narrow_margins(doc.sections[0])
    add_header_footer(doc)
    cover_page(doc)

    processed = set()
    for sec in soup.find_all("section"):
        sid = sec.get("id") or ""
        if sid in processed:
            continue
        if sid == "intent" or sid == "who":
            add_generic_section(doc, sec)
            processed.add(sid)
        elif sid == "year-map":
            add_year_map(doc, soup)
            processed.add(sid)
        elif sid in {"term1", "term2", "term3"}:
            h2 = sec.select_one("h2")
            en = sec.select_one(".en")
            banner = sec.select_one(".term-banner")
            intro = banner.find("p") if banner else None
            sub = ""
            if en:
                sub = en.get_text(" ", strip=True)
            if intro:
                sub = (sub + " — " if sub else "") + intro.get_text(" ", strip=True)
            banner_table(
                doc,
                h2.get_text(" ", strip=True) if h2 else sid,
                sub,
                fill="541024",
            )
            for article in sec.select("article.lesson"):
                add_lesson(doc, article)
            processed.add(sid)
        else:
            add_generic_section(doc, sec)
            processed.add(sid)

    # closing
    p = para(doc, space_before=18)
    add_run(p, "End of Class 3 study plan  ·  The London Tamil Sangam Tamil Language School", name=FONT_SERIF, size=11, italic=True, color=SOFT)
    p = para(doc)
    add_run(
        p,
        "This plan is based on the teacher handbook. Page numbers refer to the Class 3 textbook. "
        "Teachers should use their own craft to bring each lesson to life.",
        size=10,
        color=SOFT,
    )

    doc.save(OUT_DOCX)
    print("Wrote", OUT_DOCX, "bytes", OUT_DOCX.stat().st_size)


if __name__ == "__main__":
    main()
